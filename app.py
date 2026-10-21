import streamlit as st
import numpy as np
import pandas as pd
import matplotlib.pyplot as plt
import scipy.stats as stats
import re
import io
from docx import Document
from docx.shared import Inches

def sanitize_filename(name):
    name = re.sub(r'[\\/:*?"<>|]+', '_', str(name)).strip()
    return name or "Distribution_Capability_Report"

# ==========================================
# 0. SPC GRADING FUNCTIONS (English)
# ==========================================
def get_ca_grade(ca):
    abs_ca = abs(ca)
    if abs_ca <= 12.5: return 'A (Good)'
    elif abs_ca <= 25.0: return 'B (Fair)'
    elif abs_ca <= 50.0: return 'C (Poor)'
    else: return 'D (Inadequate)'

def get_cp_grade(cp):
    if cp >= 1.33: return 'A (Good)'
    elif cp >= 1.00: return 'B (Fair)'
    elif cp >= 0.67: return 'C (Poor)'
    else: return 'D (Inadequate)'

def get_cpk_grade(cpk):
    if cpk >= 1.67: return 'A+ (Excellent)'
    elif cpk >= 1.33: return 'A (Good)'
    elif cpk >= 1.00: return 'B (Fair)'
    elif cpk >= 0.67: return 'C (Poor)'
    else: return 'D (Inadequate)'

# ==========================================
# 0. WORD EXPORT FUNCTION
# ==========================================
def generate_word_report(title, df_specs, fig, df_metrics, df_raw):
    doc = Document()
    doc.add_heading(title, 0)

    # 1. Specifications Table
    doc.add_heading('1. Specifications', level=1)
    t_specs = doc.add_table(rows=1, cols=len(df_specs.columns))
    t_specs.style = 'Table Grid'
    for i, col_name in enumerate(df_specs.columns):
        t_specs.cell(0, i).text = str(col_name)
    for row in df_specs.itertuples(index=False):
        row_cells = t_specs.add_row().cells
        for i, val in enumerate(row):
            row_cells[i].text = str(val)

    # 2. Chart
    doc.add_heading('2. Distribution State Chart', level=1)
    img_buf = io.BytesIO()
    fig.savefig(img_buf, format='png', bbox_inches='tight', dpi=300)
    img_buf.seek(0)
    doc.add_picture(img_buf, width=Inches(6.5))

    # 3. Metrics Table
    doc.add_heading('3. Process Capability Summary', level=1)
    t_metrics = doc.add_table(rows=1, cols=len(df_metrics.columns))
    t_metrics.style = 'Table Grid'
    for i, col_name in enumerate(df_metrics.columns):
        t_metrics.cell(0, i).text = str(col_name)
    for row in df_metrics.itertuples(index=False):
        row_cells = t_metrics.add_row().cells
        for i, val in enumerate(row):
            row_cells[i].text = str(val)
            
    # 4. Input Data View Table
    doc.add_heading('4. Input Data View', level=1)
    t_raw = doc.add_table(rows=1, cols=len(df_raw.columns) + 1)
    t_raw.style = 'Table Grid'
    t_raw.cell(0, 0).text = "Row"
    for i, col_name in enumerate(df_raw.columns):
        t_raw.cell(0, i+1).text = str(col_name)
        
    for row_idx, row in enumerate(df_raw.itertuples(index=False)):
        row_cells = t_raw.add_row().cells
        row_cells[0].text = f"Row {row_idx + 1}"
        for i, val in enumerate(row):
            if pd.isna(val):
                row_cells[i+1].text = ""
            else:
                row_cells[i+1].text = f"{val:.4f}"
            
    doc_buf = io.BytesIO()
    doc.save(doc_buf)
    doc_buf.seek(0)
    return doc_buf

# ==========================================
# 1. STREAMLIT UI SETUP & SESSION STATE
# ==========================================
st.set_page_config(layout="wide", page_title="Process Capability & Distribution")

if "generate_clicked" not in st.session_state:
    st.session_state.generate_clicked = False
if "needs_update" not in st.session_state:
    st.session_state.needs_update = False

st.sidebar.header("Process Parameters")
report_title = st.sidebar.text_input("Report Title", value="Process Capability Analysis Report")

# --- TRỤC X TÙY CHỈNH TÊN ---
x_axis_label = st.sidebar.text_input("X-Axis Label (Tên trục X)", value="Measured Thickness (mm)")

st.sidebar.markdown("**Primary Standard**")
primary_standard_name = st.sidebar.text_input(
    "Primary Standard Name",
    value="YP"
)
lsl = st.sidebar.number_input(
    "Primary LSL",
    value=0.330,
    format="%.3f"
)
usl = st.sidebar.number_input(
    "Primary USL",
    value=0.376,
    format="%.3f"
)

st.sidebar.markdown("**Secondary / Reference Standard**")
secondary_standard_name = st.sidebar.text_input(
    "Secondary Standard Name",
    value="Reference Standard"
)
secondary_lsl = st.sidebar.number_input(
    "Secondary LSL",
    value=0.320,
    format="%.3f"
)
secondary_usl = st.sidebar.number_input(
    "Secondary USL",
    value=0.390,
    format="%.3f"
)

st.sidebar.markdown("---")
st.sidebar.header("Data Source")
data_source = st.sidebar.radio("Input Method:", ("Auto Generate (AI Target)", "Manual Data Entry"))

if data_source == "Auto Generate (AI Target)":
    sample_size = st.sidebar.number_input("Total Sample Size", value=75, step=1)
    st.sidebar.subheader("Target Capability (Grade A or higher)")
    target_cp_min = st.sidebar.number_input("Min Target Cp/Cpk", value=1.35) 
    target_cp_max = st.sidebar.number_input("Max Target Cp/Cpk", value=1.91)
else:
    st.sidebar.subheader("Paste Data")
    manual_input = st.sidebar.text_area("Raw Data Input (from Excel)", height=200)

if st.sidebar.button("Generate Reports"):
    st.session_state.generate_clicked = True
    st.session_state.needs_update = True

# ==========================================
# 2. XỬ LÝ DỮ LIỆU
# ==========================================
if st.session_state.generate_clicked:
    
    if usl <= lsl or secondary_usl <= secondary_lsl:
        st.error("Input Error: USL must be greater than LSL.")
        st.stop()
        
    center = (usl + lsl) / 2
    tolerance = usl - lsl
    
    if st.session_state.needs_update:
        raw_data = None
        
        # --- DATA GENERATION (AUTO) ---
        if data_source == "Auto Generate (AI Target)":
            max_attempts = 2000
            for attempt in range(max_attempts):
                t_cp = np.random.uniform(target_cp_min, target_cp_max)
                t_sigma = tolerance / (6 * t_cp)
                max_off = 0.12 * (tolerance / 2) 
                t_mean = center + np.random.choice([-1, 1]) * np.random.uniform(0.00, max_off)
                
                temp_raw = np.random.normal(loc=t_mean, scale=t_sigma, size=int(sample_size))
                temp_raw = np.round(temp_raw, 4)
                
                c_mean = np.mean(temp_raw)
                c_std = np.std(temp_raw, ddof=1)
                
                if c_std == 0: 
                    continue
                    
                cp_v = tolerance / (6 * c_std)
                cpk_v = min((usl - c_mean), (c_mean - lsl)) / (3 * c_std)
                ca_v = (abs(c_mean - center) / (tolerance / 2)) * 100
                
                if (cp_v >= 1.33) and (cpk_v >= 1.33) and (ca_v <= 12.5) and (cp_v <= target_cp_max + 0.1):
                    raw_data = temp_raw
                    break
            
            if raw_data is None:
                st.error("Goal-Seek failed to meet criteria. Try increasing sample size or adjusting target capability.")
                st.stop()
        
        # --- DATA PROCESSING (MANUAL) ---
        else:
            if not manual_input.strip():
                st.warning("Please paste some data in the text area first.")
                st.stop()
            try:
                parsed = [float(x.strip()) for x in re.split(r'[,\n\t ]+', manual_input) if x.strip()]
                raw_data = np.round(np.array(parsed), 4)
                
                if len(raw_data) < 10:
                    st.error("Insufficient data for analysis. Need at least 10 data points.")
                    st.stop()
            except ValueError:
                st.error("Format Error: Input must be numeric.")
                st.stop()
                
        st.session_state.raw_data = raw_data
        st.session_state.needs_update = False

    raw_data = st.session_state.raw_data

    # ==========================================
    # 3. TÍNH TOÁN NĂNG LỰC QUÁ TRÌNH
    # ==========================================
    raw_mean = np.mean(raw_data)
    raw_std = np.std(raw_data, ddof=1)
    raw_cp = tolerance / (6 * raw_std)
    raw_cpk = min((usl - raw_mean), (raw_mean - lsl)) / (3 * raw_std)
    raw_ca_val = (abs(raw_mean - center) / (tolerance / 2)) * 100

    # ==========================================
    # 4. HIỂN THỊ UI BÁO CÁO
    # ==========================================
    st.markdown(f"<h1 style='text-align: center; color: #2F5597;'>{report_title}</h1>", unsafe_allow_html=True)
    st.markdown("---")
    
    st.markdown("### Specifications")
    
    yp_tol = (usl - lsl) / 2
    secondary_tol = (secondary_usl - secondary_lsl) / 2
    
    df_specs = pd.DataFrame({
        f"{primary_standard_name} (+/-{yp_tol:.3f})": [
            f"{lsl:.3f} ~ {usl:.3f}"
        ],
        f"{secondary_standard_name} (+/-{secondary_tol:.3f})": [
            f"{secondary_lsl:.3f} ~ {secondary_usl:.3f}"
        ]
    })
    st.table(df_specs)
    
    # --- BIỂU ĐỒ PHÂN BỐ CẢI TIẾN ---
    fig_dist, ax_dist = plt.subplots(figsize=(11, 7))
    fig_dist.patch.set_facecolor('#ffffff')
    ax_dist.set_facecolor('#f8f9fa')
    
    plot_min = min(secondary_lsl, raw_data.min()) - (raw_std * 2)
    plot_max = max(secondary_usl, raw_data.max()) + (raw_std * 2)
    
    counts, bins, patches = ax_dist.hist(raw_data, bins=18, 
                                         color='#5B9BD5', edgecolor='white', linewidth=1.2, alpha=0.85, zorder=2)
    
    bin_width = bins[1] - bins[0]
    x_ax2 = np.linspace(plot_min, plot_max, 200)
    y_norm_scaled = stats.norm.pdf(x_ax2, raw_mean, raw_std) * len(raw_data) * bin_width
    ax_dist.plot(x_ax2, y_norm_scaled, color='#2F5597', linewidth=2.5, zorder=3)
    
    ax_dist.grid(axis='y', linestyle='--', alpha=0.6, zorder=1)
    
    # === GIẢI QUYẾT LỖI ĐÈ NHÃN BẰNG CÁCH NỚI RỘNG TRỤC Y VÀ PHÂN TẦNG (STAGGERING) ===
    y_max = ax_dist.get_ylim()[1]
    ax_dist.set_ylim(0, y_max * 1.30)  # Tăng thêm 30% khoảng không phía trên biểu đồ
    
    text_y_center = y_max * 1.15  # Tầng 1: YP Center (Cao nhất)
    text_y_top = y_max * 0.90     # Tầng 2: YP Limits (Thấp hơn hộp Stats)
    text_y_mid = y_max * 0.70     # Tầng 3: Secondary Limits
    
    bbox_center = dict(boxstyle="round,pad=0.3", fc="#e6ffe6", ec="green", alpha=0.9)
    bbox_yp = dict(boxstyle="round,pad=0.3", fc="#ffe6e6", ec="red", alpha=0.9)
    bbox_secondary = dict(boxstyle="round,pad=0.3", fc="#e6f2ff", ec="blue", alpha=0.9)
    
    # Vẽ Center
    ax_dist.axvline(center, color='green', linestyle='-', linewidth=1.5, alpha=0.8, zorder=4)
    ax_dist.text(center, text_y_center, f'{primary_standard_name} Center', color='darkgreen', 
                 ha='center', va='bottom', fontsize=10, fontweight='bold', bbox=bbox_center, zorder=5)

    # Vẽ YP Limits
    ax_dist.axvline(lsl, color='red', linestyle='--', linewidth=1.5, zorder=4)
    ax_dist.text(lsl, text_y_top, f'{primary_standard_name} LSL\nLower', color='darkred', 
                 ha='center', va='top', fontsize=9, fontweight='bold', bbox=bbox_yp, zorder=5)
    
    ax_dist.axvline(usl, color='red', linestyle='--', linewidth=1.5, zorder=4)
    ax_dist.text(usl, text_y_top, f'{primary_standard_name} USL\nUpper', color='darkred', 
                 ha='center', va='top', fontsize=9, fontweight='bold', bbox=bbox_yp, zorder=5)

    # Vẽ Secondary Limits
    ax_dist.axvline(secondary_lsl, color='blue', linestyle='-.', linewidth=1.5, zorder=4)
    ax_dist.text(secondary_lsl, text_y_mid, f'{secondary_standard_name} LSL\nLower', color='darkblue', 
                 ha='center', va='bottom', fontsize=9, fontweight='bold', bbox=bbox_secondary, zorder=5)
    
    ax_dist.axvline(secondary_usl, color='blue', linestyle='-.', linewidth=1.5, zorder=4)
    ax_dist.text(secondary_usl, text_y_mid, f'{secondary_standard_name} USL\nUpper', color='darkblue', 
                 ha='center', va='bottom', fontsize=9, fontweight='bold', bbox=bbox_secondary, zorder=5)

    # Stats Box an toàn ở góc phải trên cùng
    stats_text = f"Mean = {raw_mean:.4f}\nStd. Dev. = {raw_std:.4f}\nN = {len(raw_data)}"
    ax_dist.text(0.98, 0.96, stats_text, transform=ax_dist.transAxes,
                 fontsize=11, va='top', ha='right',
                 bbox=dict(boxstyle='square,pad=0.5', facecolor='white', edgecolor='gray', alpha=0.9), zorder=6)

    ax_dist.set_xlim(plot_min, plot_max)
    
    # --- ÁP DỤNG TRỤC X TÙY CHỈNH VÀO BIỂU ĐỒ ---
    ax_dist.set_xlabel(x_axis_label, fontsize=12, fontweight='bold', labelpad=15)
    ax_dist.set_ylabel('Number of Steel Coils', fontsize=12, fontweight='bold', labelpad=15)
    
    st.pyplot(fig_dist)
    
    st.markdown("---")
    st.markdown("### Process Capability Summary Table")
    
    df_metrics = pd.DataFrame({
        "Sample Size (N)": [len(raw_data)],
        f"{primary_standard_name} LSL": [f"{lsl:.3f}"],
        f"{primary_standard_name} USL": [f"{usl:.3f}"],
        "Mean": [f"{raw_mean:.4f}"],
        "Std Dev": [f"{raw_std:.5f}"],
        "Accuracy (Ca)": [f"{raw_ca_val:.2f}%"],
        "Precision (Cp)": [f"{raw_cp:.3f}"],
        "Capability (Cpk)": [f"{raw_cpk:.3f}"],
        "Ca Grade": [get_ca_grade(raw_ca_val)],
        "Cp Grade": [get_cp_grade(raw_cp)],
        "Cpk Grade": [get_cpk_grade(raw_cpk)]
    })
    
    st.dataframe(df_metrics, hide_index=True, use_container_width=True)

    # --- INPUT DATA VIEW ---
    st.markdown("---")
    st.markdown("### Input Data View")
    st.write(f"Total samples recorded: **{len(raw_data)}**")
    
    cols = min(10, len(raw_data))
    rows = int(np.ceil(len(raw_data) / cols))
    padded_data = np.pad(raw_data, (0, (cols * rows) - len(raw_data)), constant_values=np.nan)
    df_raw_view = pd.DataFrame(padded_data.reshape(rows, cols))
    df_raw_view.columns = [f"Col {i+1}" for i in range(cols)]
    df_raw_view.index = [f"Row {i+1}" for i in range(rows)]
    
    st.dataframe(df_raw_view.style.format("{:.4f}", na_rep=""), use_container_width=True)
    
    # --- WORD EXPORT ---
    st.markdown("<br>", unsafe_allow_html=True)
    word_file = generate_word_report(report_title, df_specs, fig_dist, df_metrics, df_raw_view)
    
    st.download_button(
        label="📥 Download Report (Word)",
        data=word_file.getvalue(),
        file_name=f"{sanitize_filename(report_title)}.docx",
        mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        type="primary"
    )
